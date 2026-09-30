"""PowerPoint (.pptx) document extractor used for RAG ingestion.

Load ``.pptx`` files into Markdown-ish text with embedded image links. Each
slide becomes a ``## Slide N`` section, tables become Markdown tables, and
pictures are persisted through the storage layer so the frontend can still
render them after the file has been vectorized.
"""

import hashlib
import logging
import mimetypes
import os
import uuid
from typing import override

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy.orm import Session

from configs import dify_config
from core.rag.extractor.extractor_base import BaseExtractor
from core.rag.models.document import Document
from extensions.ext_database import db
from extensions.ext_storage import storage
from extensions.storage.storage_type import StorageType
from libs.datetime_utils import naive_utc_now
from models.enums import CreatorUserRole
from models.model import UploadFile

logger = logging.getLogger(__name__)


class PptxExtractor(BaseExtractor):
    """Load ``.pptx`` files into Markdown-ish text with embedded image links.

    Args:
        file_path: Path to the file to load.
        tenant_id: Tenant that owns the extracted images.
        user_id: Account used as the creator of the extracted images.
        session: Session used to persist extracted images.
    """

    _closed: bool
    _session: Session | None

    def __init__(self, file_path: str, tenant_id: str, user_id: str, *, session: Session | None = None):
        """Initialize with file path."""
        self._closed = False
        self.file_path = file_path
        self.tenant_id = tenant_id
        self.user_id = user_id
        self._session = session

        if "~" in self.file_path:
            self.file_path = os.path.expanduser(self.file_path)

        if not os.path.isfile(self.file_path):
            raise ValueError(f"File path {self.file_path} is not a valid file")

    @override
    def extract(self) -> list[Document]:
        content = self.parse_pptx(self.file_path)
        return [Document(page_content=content, metadata={"source": self.file_path})]

    @staticmethod
    def _image_key(blob: bytes) -> str:
        """Content hash used to deduplicate identical images."""
        return hashlib.sha1(blob).hexdigest()

    @staticmethod
    def _shape_sort_key(shape) -> tuple[float, float]:
        """Order shapes by reading position: top first, then left."""
        try:
            return (shape.top or 0, shape.left or 0)
        except Exception:
            return (0, 0)

    def _save_image(self, blob: bytes, ext: str) -> str:
        """Persist one image and return its Markdown link."""
        ext = (ext or "png").lstrip(".").lower()
        if ext == "jpg":
            ext = "jpeg"

        file_uuid = str(uuid.uuid4())
        file_key = "image_files/" + self.tenant_id + "/" + file_uuid + "." + ext
        mime_type, _ = mimetypes.guess_type(file_key)

        storage.save(file_key, blob)

        upload_file = UploadFile(
            tenant_id=self.tenant_id,
            storage_type=StorageType(dify_config.STORAGE_TYPE),
            key=file_key,
            name=file_key,
            size=len(blob),
            extension=ext,
            mime_type=mime_type or "",
            created_by=self.user_id,
            created_by_role=CreatorUserRole.ACCOUNT,
            created_at=naive_utc_now(),
            used=True,
            used_by=self.user_id,
            used_at=naive_utc_now(),
        )
        session = self._session or db.session
        session.add(upload_file)

        base_url = dify_config.FILES_URL
        link = f"![image]({base_url}/files/{upload_file.id}/file-preview)"
        logger.info("PptxExtractor saved image key=%s id=%s", file_key, upload_file.id)
        return link

    def _table_to_markdown(self, table) -> str:
        """Render a PowerPoint table as a Markdown table."""
        lines: list[str] = []
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if row_idx == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        return "\n".join(lines)

    def _process_shape(self, shape, out: list[str], image_cache: dict[str, str]) -> None:
        """Append the Markdown representation of one shape to ``out``."""
        try:
            shape_type = shape.shape_type
        except Exception:
            shape_type = None

        # 1) picture -> deduplicated Markdown image link
        if shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                image = shape.image
                blob = image.blob
            except Exception:
                logger.warning("PptxExtractor skipped an unreadable picture")
                return

            key = self._image_key(blob)
            link = image_cache.get(key)
            if link is None:
                try:
                    link = self._save_image(blob, getattr(image, "ext", "png"))
                except Exception:
                    logger.exception("PptxExtractor failed to persist an image")
                    return
                image_cache[key] = link
            out.append(link)
            return

        # 2) table -> Markdown table
        if getattr(shape, "has_table", False):
            try:
                markdown = self._table_to_markdown(shape.table)
            except Exception:
                logger.warning("PptxExtractor failed to render a table")
                markdown = ""
            if markdown:
                out.append(markdown)
            return

        # 3) grouped shapes -> walk children
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                for child in sorted(shape.shapes, key=self._shape_sort_key):
                    self._process_shape(child, out, image_cache)
            except Exception:
                logger.warning("PptxExtractor failed to walk a group shape")
            return

        # 4) text frame -> raw text
        if getattr(shape, "has_text_frame", False):
            try:
                text = (shape.text_frame.text or "").strip()
            except Exception:
                text = ""
            if text:
                out.append(text)
            return

        # 5) charts, embedded objects and other non-textual elements
        out.append("<!-- non-textual element -->")

    def parse_pptx(self, pptx_path: str) -> str:
        prs = Presentation(pptx_path)
        image_cache: dict[str, str] = {}
        slides: list[str] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in sorted(slide.shapes, key=self._shape_sort_key):
                self._process_shape(shape, parts, image_cache)
            body = "\n".join(p for p in parts if p and p.strip())
            slides.append(f"## Slide {slide_idx}\n{body}")

        content = "\n\n".join(slides)
        if self._session is None:
            db.session.commit()

        logger.info(
            "PptxExtractor parsed %s: %d slides, %d unique images saved",
            self.file_path,
            len(slides),
            len(image_cache),
        )
        return content
